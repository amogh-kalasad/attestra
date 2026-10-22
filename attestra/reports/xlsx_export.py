"""Engagement XLSX export with Findings, Controls, and Evidence metadata sheets.

Evidence bodies and bytes are never included. Formula-injection guarding reuses
csv_cell from findings_csv.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from attestra import audit, services
from attestra.models import Evidence, Finding
from attestra.reports.findings_csv import CSV_COLUMNS, csv_cell, findings_to_rows

CONTROLS_COLUMNS = (
    "framework",
    "ref",
    "title",
    "status",
    "proposed",
    "reason",
)

EVIDENCE_COLUMNS = (
    "finding_id",
    "evidence id",
    "kind",
    "filename",
    "sha256",
    "note",
)


def _write_sheet(ws, headers: tuple[str, ...] | list[str], rows: list[list[str]]) -> None:
    """Write header + data, freeze the top row, and enable auto-filter."""
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _controls_rows(s: Session, engagement_id: int) -> list[list[str]]:
    eng = services.get_engagement(s, engagement_id)
    rows: list[list[str]] = []
    for framework in eng.frameworks or []:
        for r in services.compliance(s, engagement_id, framework):
            rows.append(
                [
                    csv_cell(framework),
                    csv_cell(r.ref),
                    csv_cell(r.title),
                    csv_cell(r.status),
                    csv_cell(r.proposed),
                    csv_cell(r.reason),
                ]
            )
    return rows


def _evidence_rows(s: Session, engagement_id: int) -> list[list[str]]:
    """Metadata only — never path, body bytes, or masking details."""
    evidence = list(
        s.scalars(
            select(Evidence)
            .join(Finding, Evidence.finding_id == Finding.id)
            .where(Finding.engagement_id == engagement_id)
            .order_by(Evidence.id)
        )
    )
    return [
        [
            csv_cell(ev.finding_id),
            csv_cell(ev.id),
            csv_cell(ev.kind),
            csv_cell(ev.filename),
            csv_cell(ev.sha256),
            csv_cell(ev.note),
        ]
        for ev in evidence
    ]


def render_engagement_xlsx(s: Session, engagement_id: int) -> bytes:
    """Build the three-sheet workbook bytes without auditing."""
    findings = services.list_findings(s, engagement_id)
    wb = Workbook()
    ws_findings = wb.active
    ws_findings.title = "Findings"
    _write_sheet(ws_findings, CSV_COLUMNS, findings_to_rows(findings))

    ws_controls = wb.create_sheet("Controls")
    _write_sheet(ws_controls, CONTROLS_COLUMNS, _controls_rows(s, engagement_id))

    ws_evidence = wb.create_sheet("Evidence")
    _write_sheet(ws_evidence, EVIDENCE_COLUMNS, _evidence_rows(s, engagement_id))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_engagement_xlsx(s: Session, engagement_id: int, *, actor: str) -> bytes:
    """Export Findings/Controls/Evidence sheets for an engagement and audit it."""
    services.get_engagement(s, engagement_id)
    body = render_engagement_xlsx(s, engagement_id)
    audit.record(
        s,
        actor,
        "engagement.exported",
        "engagement",
        engagement_id,
        format="xlsx",
    )
    return body
