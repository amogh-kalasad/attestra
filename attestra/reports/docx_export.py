"""Production-styled Word (.docx) export built with python-docx.

The frozen HTML report is parsed and rebuilt as native Word content:

* cover page (no header/footer on it) followed by a Contents page with a real TOC field;
* real ``Heading 1``/``Heading 2``/``Heading 3`` styles, so Word's navigation pane and TOC work;
* tables with a shaded, repeating header row, sized columns and consistent cell padding;
* severity and status badges as centred, colour-shaded cells (or shaded runs inline);
* status/KPI tiles as a single spaced table row;
* header with tenant and report title, footer with ``Page X of Y``;
* a DRAFT watermark and a red header band until the report is signed;
* one font family (Calibri, Consolas for code) and a fixed size scale;
* Title Case display labels with security acronyms preserved.

Only displayed text is changed. Stored values, hashes and the API are untouched.
"""
# @lat: [[reports#Exports]]

from __future__ import annotations

import base64
import re
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

FONT = "Calibri"
MONO = "Consolas"
NAVY = "1F3A5F"
INK = "1F2937"
MUTED = "6B7280"
GRID = "C9D1DC"
HEAD_FILL = "1F3A5F"
KEY_FILL = "EEF2F7"
ZEBRA = "F7F9FC"

SEV_COLOURS = {"critical": "A1161B", "high": "C2410C", "medium": "A16207", "low": "15803D", "info": "475569"}
ST_COLOURS = {
    "met": "15803D", "no_exceptions": "0F766E", "partial": "A16207", "not_met": "A1161B",
    "not_assessed": "64748B", "not_applicable": "64748B", "fixed": "15803D", "still_open": "A1161B",
}
LETTER_KINDS = {"closure_letter", "attestation_letter"}

ACRONYMS = {
    "pci", "dss", "cvss", "owasp", "vapt", "va", "pt", "iso", "iec", "soc", "nist", "csf", "rbi", "sebi",
    "cscrf", "dora", "cert-in", "cwe", "cve", "epss", "kev", "asvs", "wstg", "tls", "ssl", "sla", "sha-256",
    "sha", "url", "api", "id", "ip", "dns", "http", "https", "csp", "xss", "sql", "qa", "ai", "it", "eu",
    "cde", "pan", "soa", "ui", "pdf", "csv", "sarif", "oscal", "json", "hsts", "ocsp", "tsc", "dast", "sast",
    "pii", "mfa", "sso", "ca", "eng", "it", "f",
}
SMALL = {"a", "an", "and", "as", "at", "but", "by", "for", "in", "of", "on", "or", "per", "the", "to", "vs", "via", "with"}

BLOCK = {"p", "div", "section", "h1", "h2", "h3", "h4", "h5", "h6", "table", "ul", "ol", "li", "nav", "footer",
         "header", "figure", "pre", "blockquote", "hr", "main", "article", "svg"}
VOID = {"br", "img", "meta", "link", "hr", "input", "col", "source", "wbr", "rect"}
SKIP = {"head", "style", "script", "title", "button", "form", "select", "noscript", "template"}


# --------------------------------------------------------------------------- labels

def title_case(text: str) -> str:
    """Title Case a display label, keeping acronyms (PCI DSS, CVSS ...) and identifiers intact."""
    words = re.split(r"(\s+)", text.replace("_", " "))
    out, first = [], True
    for w in words:
        if not w or w.isspace():
            out.append(w)
            continue
        core = re.sub(r"^[^\w]+|[^\w]+$", "", w)
        low = core.lower()
        if not core or re.search(r"\d", core) and not re.fullmatch(r"[a-z]+", low) or "." in core or "/" in w and len(core) > 3 and low not in ACRONYMS:
            # numbers, refs (1.2.1, F-73), paths, urls: leave alone
            out.append(w)
        elif low in ACRONYMS and (core.islower() or core.isupper()):
            out.append(w.replace(core, core.upper()))
        elif any(c.isupper() for c in core[1:]) or core.isupper():
            out.append(w)  # already mixed / upper case (HarmonyPM, ATTESTRA)
        elif low in SMALL and not first:
            out.append(w.replace(core, low))
        else:
            out.append(w.replace(core, core[:1].upper() + core[1:]))
        first = False
    return "".join(out)


def is_label(text: str) -> bool:
    """A short lowercase token-ish value that is a label, not prose or data."""
    t = text.strip()
    return bool(t) and len(t) <= 40 and len(t.split()) <= 4 and bool(re.fullmatch(r"[a-z][a-z _/()-]*", t))


FRAMEWORK_NAMES = {
    "pci_dss_4": "PCI DSS v4.0.1", "iso_27001_2022": "ISO/IEC 27001:2022", "soc2": "SOC 2",
    "iso_42001_2023": "ISO/IEC 42001:2023", "cert_in_2022": "CERT-In 2022", "nist_csf_2": "NIST CSF 2.0",
    "rbi_csf_2016": "RBI CSF 2016", "sebi_cscrf_2024": "SEBI CSCRF 2024", "dora_2022": "DORA",
}
KIND_NAMES = {"closure_letter": "Closure Letter", "attestation_letter": "Attestation Letter", "iso_soa": "ISO SoA"}
_TS = re.compile(r"(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2})(?::\d{2}(?:\.\d+)?)?(?:\+00:00|Z)")


def display(text: str) -> str:
    """Human-readable display text: raw framework ids, enum kinds and ISO timestamps (display only)."""
    text = _TS.sub(lambda m: f"{m.group(1)} {m.group(2)} UTC", text)
    text = re.sub(r"\b(" + "|".join(FRAMEWORK_NAMES) + r")\b", lambda m: FRAMEWORK_NAMES[m.group(1)], text)
    text = re.sub(r"\b(closure_letter|attestation_letter|iso_soa)\b", lambda m: KIND_NAMES[m.group(1).lower()], text,
                  flags=re.I)
    return text


# --------------------------------------------------------------------------- html tree

class Node:
    __slots__ = ("tag", "attrs", "children")

    def __init__(self, tag: str, attrs: dict | None = None):
        self.tag, self.attrs, self.children = tag, attrs or {}, []

    @property
    def classes(self) -> set[str]:
        return set((self.attrs.get("class") or "").split())

    def find_all(self, tag: str):
        for c in self.children:
            if isinstance(c, Node):
                if c.tag == tag:
                    yield c
                yield from c.find_all(tag)


class _Tree(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = Node(tag, {k: (v or "") for k, v in attrs})
        self.stack[-1].children.append(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].children.append(Node(tag, {k: (v or "") for k, v in attrs}))

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        self.stack[-1].children.append(data)


def parse(html: str) -> Node:
    t = _Tree()
    t.feed(html)
    return t.root


def text_of(node, keep_ws: bool = False) -> str:
    if isinstance(node, str):
        return node
    if node.tag in SKIP:
        return ""
    if node.tag == "br":
        return "\n"
    s = "".join(text_of(c, keep_ws) for c in node.children)
    return s if keep_ws else s


def norm(s: str) -> str:
    return re.sub(r"[ \t\r\n\xa0]+", " ", s).strip()


def badge_of(node) -> tuple[str, str] | None:
    """Return (label, hex colour) if node is a severity/status badge span."""
    if isinstance(node, str) or node.tag != "span":
        return None
    cls = node.classes
    if "sev" in cls or "st" in cls:
        colour = "475569"
        for c in cls:
            if c.startswith("sev-"):
                colour = SEV_COLOURS.get(c[4:], colour)
            elif c.startswith("st-"):
                colour = ST_COLOURS.get(c[3:], colour)
        m = re.search(r"color:\s*#([0-9a-fA-F]{6})", node.attrs.get("style", ""))
        bg = re.search(r"background:\s*#([0-9a-fA-F]{6})", node.attrs.get("style", ""))
        if bg and m:
            colour = m.group(1)  # tag-style chip: use its ink colour as the fill
        return title_case(norm(text_of(node))), colour.upper()
    return None


# --------------------------------------------------------------------------- docx helpers

def _qn(tag):
    from docx.oxml.ns import qn
    return qn(tag)


def _el(tag, **attrs):
    from docx.oxml import OxmlElement
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(_qn(k.replace("__", ":")), str(v))
    return e


def shade(cell_or_ppr, fill: str):
    tcPr = cell_or_ppr._tc.get_or_add_tcPr() if hasattr(cell_or_ppr, "_tc") else cell_or_ppr
    for old in tcPr.findall(_qn("w:shd")):
        tcPr.remove(old)
    tcPr.append(_el("w:shd", **{"w__val": "clear", "w__color": "auto", "w__fill": fill}))


def run_shade(run, fill: str):
    rPr = run._r.get_or_add_rPr()
    rPr.append(_el("w:shd", **{"w__val": "clear", "w__color": "auto", "w__fill": fill}))


def cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    mar = _el("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        mar.append(_el(f"w:{side}", **{"w__w": v, "w__type": "dxa"}))
    tblPr.append(mar)


def table_borders(table, colour=GRID, size=4, inside=True):
    tblPr = table._tbl.tblPr
    b = _el("w:tblBorders")
    sides = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for s in sides:
        b.append(_el(f"w:{s}", **{"w__val": "single", "w__sz": size, "w__space": 0, "w__color": colour}))
    if not inside:
        for s in ("insideH", "insideV"):
            b.append(_el(f"w:{s}", **{"w__val": "nil"}))
    tblPr.append(b)


def no_borders(table):
    tblPr = table._tbl.tblPr
    b = _el("w:tblBorders")
    for s in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(_el(f"w:{s}", **{"w__val": "nil"}))
    tblPr.append(b)


def fixed_layout(table, widths_twips: list[int]):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblPr.append(_el("w:tblLayout", **{"w__type": "fixed"}))
    tblW = tblPr.find(_qn("w:tblW"))
    if tblW is None:
        tblW = _el("w:tblW")
        tblPr.append(tblW)
    tblW.set(_qn("w:w"), str(sum(widths_twips)))
    tblW.set(_qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(_qn("w:gridCol")), widths_twips):
        gc.set(_qn("w:w"), str(w))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_twips):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(_qn("w:tcW"))
            if tcW is None:
                tcW = _el("w:tcW")
                tcPr.insert(0, tcW)
            tcW.set(_qn("w:w"), str(w))
            tcW.set(_qn("w:type"), "dxa")


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader", **{"w__val": "true"}))


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:cantSplit", **{"w__val": "true"}))


def v_align(cell, val="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(_el("w:vAlign", **{"w__val": val}))


def add_field(paragraph, code: str, placeholder: str = "1"):
    r = paragraph.add_run()
    r._r.append(_el("w:fldChar", **{"w__fldCharType": "begin"}))
    r2 = paragraph.add_run()
    it = _el("w:instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    it.text = f" {code} "
    r2._r.append(it)
    r3 = paragraph.add_run()
    r3._r.append(_el("w:fldChar", **{"w__fldCharType": "separate"}))
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run()
    r5._r.append(_el("w:fldChar", **{"w__fldCharType": "end"}))
    return [r, r2, r3, r4, r5]


def para_border_bottom(paragraph, colour=NAVY, size=8):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = _el("w:pBdr")
    bdr.append(_el("w:bottom", **{"w__val": "single", "w__sz": size, "w__space": 4, "w__color": colour}))
    pPr.append(bdr)


def para_shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(_el("w:shd", **{"w__val": "clear", "w__color": "auto", "w__fill": fill}))


def keep_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


WATERMARK_XML = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
    '<w:pict><v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" '
    'path="m@7,l@8,m@5,21600l@6,21600e"><v:path textpathok="t" o:connecttype="custom"/>'
    '<v:textpath on="t" fitshape="t"/></v:shapetype>'
    '<v:shape id="AttestraDraftWatermark" o:spid="_x0000_s2049" type="#_x0000_t136" '
    'style="position:absolute;margin-left:0;margin-top:0;width:430pt;height:140pt;rotation:315;z-index:-251654144;'
    'mso-position-horizontal:center;mso-position-horizontal-relative:margin;'
    'mso-position-vertical:center;mso-position-vertical-relative:margin" fillcolor="#E5E7EB" stroked="f">'
    '<v:fill opacity=".45"/><v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt;font-weight:bold" '
    'string="DRAFT"/></v:shape></w:pict></w:r>'
)


# --------------------------------------------------------------------------- builder

class Builder:
    def __init__(self, *, kind: str, draft: bool, base_url: str | None):
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Cm, Pt, RGBColor

        self.Pt, self.Cm, self.RGB = Pt, Cm, RGBColor
        self.kind, self.draft = kind, draft
        self.base = Path(base_url) if base_url else None
        self.doc = Document()
        self.headings: list[tuple[int, str]] = []
        self.h_offset = 0
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
        sec.header_distance = sec.footer_distance = Cm(1.0)
        self.text_w = int((21.0 - 4.0) / 2.54 * 1440)  # twips

        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name, normal.font.size = FONT, Pt(10)
        normal.font.color.rgb = RGBColor.from_string(INK)
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(_qn("w:rFonts"))
        if rfonts is None:
            rfonts = _el("w:rFonts")
            rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(_qn(a), FONT)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, colour, before, after in (
            ("Title", 26, NAVY, 0, 6), ("Heading 1", 16, NAVY, 18, 6),
            ("Heading 2", 13, NAVY, 12, 4), ("Heading 3", 11, INK, 10, 3),
        ):
            s = st[name]
            s.font.name, s.font.size, s.font.bold = FONT, Pt(size), True
            s.font.color.rgb = RGBColor.from_string(colour)
            r = s.element.get_or_add_rPr()
            rf = r.find(_qn("w:rFonts"))
            if rf is None:
                rf = _el("w:rFonts")
                r.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(_qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rf.get(_qn(a)) is not None:
                    del rf.attrib[_qn(a)]
            s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
            s.paragraph_format.keep_with_next = True
        title_pf = st["Title"].element.get_or_add_pPr()
        for b in title_pf.findall(_qn("w:pBdr")):
            title_pf.remove(b)
        for name, size, colour, mono in (("Meta", 9, MUTED, False), ("Code", 8.5, INK, True), ("Small", 8, MUTED, False)):
            s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = normal
            s.font.size, s.font.color.rgb = Pt(size), RGBColor.from_string(colour)
            if mono:
                s.font.name = MONO
                rp = s.element.get_or_add_rPr()
                rf = _el("w:rFonts", **{"w__ascii": MONO, "w__hAnsi": MONO, "w__cs": MONO})
                rp.append(rf)
                s.paragraph_format.space_after = Pt(0)
                s.paragraph_format.line_spacing = 1.0
        for lvl in (1, 2, 3):
            try:
                ts = st[f"TOC {lvl}"]
            except KeyError:
                ts = st.add_style(f"TOC {lvl}", WD_STYLE_TYPE.PARAGRAPH)
                ts.base_style = normal
            ts.font.name, ts.font.size = FONT, Pt(10.5 if lvl == 1 else 10)
            ts.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
            ts.paragraph_format.space_after = Pt(3)

    # ---- inline runs
    def runs(self, par, node, fmt=None, label=False):
        fmt = dict(fmt or {})
        if isinstance(node, str):
            t = re.sub(r"[ \t\r\n]+", " ", node) if not fmt.get("pre") else node
            if not t:
                return
            if label and is_label(t):
                t = title_case(t)
            self._run(par, t, fmt)
            return
        if node.tag in SKIP:
            return
        if node.tag == "br":
            par.add_run().add_break()
            return
        if node.tag == "img":
            self.image(par, node)
            return
        b = badge_of(node)
        if b:
            lab, colour = b
            r = par.add_run(f"\u00a0{lab}\u00a0")
            r.bold, r.font.size = True, self.Pt(8.5)
            r.font.color.rgb = self.RGB.from_string("FFFFFF")
            run_shade(r, colour)
            return
        if node.tag in ("b", "strong", "th"):
            fmt["bold"] = True
        if node.tag in ("i", "em"):
            fmt["italic"] = True
        if "mono" in node.classes or node.tag in ("code", "pre", "kbd"):
            fmt["mono"] = True
        if "meta" in node.classes:
            fmt["muted"] = True
            label = True
        if "brand" in node.classes:
            fmt["bold"], fmt["colour"] = True, NAVY
        m = re.search(r"color:\s*#([0-9a-fA-F]{6})", node.attrs.get("style", ""))
        if m:
            fmt["colour"] = m.group(1)
        if "font-weight" in node.attrs.get("style", ""):
            fmt["bold"] = True
        if node.tag == "a" and node.attrs.get("href", "").startswith("http"):
            fmt["colour"] = "1D4ED8"
        for c in node.children:
            self.runs(par, c, fmt, label)

    def _run(self, par, t, fmt):
        if not fmt.get("mono") and not fmt.get("pre"):
            t = display(t)
        r = par.add_run(t)
        if fmt.get("bold"):
            r.bold = True
        if fmt.get("italic"):
            r.italic = True
        if fmt.get("mono"):
            r.font.name = MONO
            r._r.get_or_add_rPr().get_or_add_rFonts().set(_qn("w:hAnsi"), MONO)
            r.font.size = self.Pt(8.5)
        if fmt.get("muted"):
            r.font.color.rgb = self.RGB.from_string(MUTED)
        if fmt.get("colour"):
            r.font.color.rgb = self.RGB.from_string(fmt["colour"])
        if fmt.get("size"):
            r.font.size = self.Pt(fmt["size"])
        if fmt.get("white"):
            r.font.color.rgb = self.RGB.from_string("FFFFFF")

    def image(self, par, node):
        src = node.attrs.get("src", "")
        data = None
        try:
            if src.startswith("data:"):
                data = base64.b64decode(src.split(",", 1)[1])
            elif src and self.base is not None and not src.startswith(("http:", "https:")):
                p = (self.base / src).resolve()
                if p.is_file():
                    data = p.read_bytes()
        except Exception:
            data = None
        if not data:
            alt = node.attrs.get("alt") or src.rsplit("/", 1)[-1]
            if alt:
                self._run(par, f"[Image: {alt}]", {"muted": True})
            return
        try:
            from docx.shared import Cm
            par.add_run().add_picture(BytesIO(data), width=Cm(15.5))
            par.alignment = 1
        except Exception:
            self._run(par, f"[Image: {node.attrs.get('alt') or 'evidence'}]", {"muted": True})

    # ---- blocks
    def para(self, style=None):
        p = self.doc.add_paragraph(style=style)
        return p

    def heading(self, node, level):
        txt = norm(text_of(node))
        if not txt:
            return
        txt = display(txt)
        if not re.match(r"^F-\d+", txt):
            txt = title_case(txt)
        level = max(1, min(3, level))
        h = self.doc.add_heading(txt, level=level)
        self.headings.append((level, txt))
        return h

    def blocks(self, node, ctx=None):
        """Render children of node as block content."""
        inline: list = []

        def flush():
            if any(not isinstance(i, str) or i.strip() for i in inline):
                p = self.para()
                for i in inline:
                    self.runs(p, i)
                if not p.text.strip() and not p._p.findall(".//" + _qn("w:drawing")):
                    p._p.getparent().remove(p._p)
            inline.clear()

        for c in node.children:
            if isinstance(c, str) or (c.tag not in BLOCK and "kpis" not in c.classes and c.tag != "table"):
                inline.append(c)
                continue
            flush()
            self.block(c)
        flush()

    def block(self, n):
        tag, cls = n.tag, n.classes
        if tag in SKIP or "draft-mark" in cls:
            return
        if tag == "nav" and "toc" in cls:
            return
        if tag == "section" and "cover" in cls:
            return  # handled by cover()
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.heading(n, int(tag[1]) - self.h_offset)
            return
        if tag == "table":
            self.table(n)
            return
        if tag == "img":
            self.image(self.para(), n)
            return
        if "kpis" in cls:
            self.tiles(n)
            return
        if tag == "svg":
            self.chart(n)
            return
        if tag in ("ul", "ol"):
            for i, li in enumerate([c for c in n.children if isinstance(c, Node) and c.tag == "li"], 1):
                p = self.para()
                pf = p.paragraph_format
                pf.left_indent, pf.first_line_indent = self.Cm(0.9), self.Cm(-0.5)
                pf.space_after = self.Pt(2)
                p.add_run(f"{i}.\t" if tag == "ol" else "\u2022\t")
                pf.tab_stops.add_tab_stop(self.Cm(0.9))
                for c in li.children:
                    if isinstance(c, Node) and c.tag in ("ul", "ol", "table"):
                        self.block(c)
                    else:
                        self.runs(p, c)
            return
        if tag == "pre":
            txt = text_of(n).strip("\n")
            tbl = self.doc.add_table(rows=1, cols=1)
            table_borders(tbl, colour="D1D5DB")
            fixed_layout(tbl, [self.text_w])
            cell_margins(tbl, 80, 80, 120, 120)
            cell = tbl.cell(0, 0)
            shade(cell, "F3F4F6")
            p = cell.paragraphs[0]
            p.style = self.doc.styles["Code"]
            lines = txt.split("\n")
            for i, line in enumerate(lines):
                p.add_run(line)
                if i < len(lines) - 1:
                    p.add_run().add_break()
            self.spacer(4)
            return
        if tag == "footer":
            p = self.para("Small")
            para_border_bottom(p, colour="FFFFFF", size=0)
            p.paragraph_format.space_before = self.Pt(14)
            for c in n.children:
                self.runs(p, c)
            return
        if "note" in cls:
            tbl = self.doc.add_table(rows=1, cols=1)
            no_borders(tbl)
            fixed_layout(tbl, [self.text_w])
            cell_margins(tbl, 100, 100, 160, 120)
            cell = tbl.cell(0, 0)
            shade(cell, "EFF6FF")
            tcPr = cell._tc.get_or_add_tcPr()
            bd = _el("w:tcBorders")
            bd.append(_el("w:left", **{"w__val": "single", "w__sz": 24, "w__space": 0, "w__color": "2563EB"}))
            tcPr.append(bd)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = self.Pt(0)
            for c in n.children:
                self.runs(p, c)
            self.spacer(6)
            return
        if "finding" in cls:
            self.finding(n)
            return
        if tag in ("p", "li") or (tag == "div" and not any(isinstance(c, Node) and (c.tag in BLOCK or c.tag == "table") for c in n.children)):
            style = "Meta" if "meta" in cls else None
            p = self.para(style)
            fmt = {"pre": True} if "pre-wrap" in n.attrs.get("style", "") else {}
            if fmt:
                txt = text_of(n).strip("\n")
                for i, line in enumerate(txt.split("\n")):
                    if i:
                        p.add_run().add_break()
                    p.add_run(line)
            else:
                for c in n.children:
                    self.runs(p, c, label="meta" in cls)
            if not p.text.strip() and not p._p.findall(".//" + _qn("w:drawing")):
                p._p.getparent().remove(p._p)
            return
        self.blocks(n)

    def spacer(self, pts=6):
        p = self.para()
        pf = p.paragraph_format
        pf.space_after, pf.space_before = self.Pt(0), self.Pt(0)
        pf.line_spacing = self.Pt(pts)

    def finding(self, n):
        for c in n.children:
            if isinstance(c, Node) and c.tag == "div" and any(badge_of(x) for x in c.children):
                self.badge_line(c)
            elif isinstance(c, Node):
                self.block(c)
            elif c.strip():
                p = self.para()
                self.runs(p, c)
        self.spacer(8)

    def badge_line(self, n):
        """Finding header: badges as shaded cells, then the status text, in one aligned row."""
        badges = [badge_of(c) for c in n.children if badge_of(c)]
        rest = [c for c in n.children if not badge_of(c)]
        widths = [1500] * len(badges)
        widths.append(self.text_w - sum(widths) - 120 * len(badges))
        tbl = self.doc.add_table(rows=1, cols=len(widths))
        no_borders(tbl)
        tblPr = tbl._tbl.tblPr
        tblPr.append(_el("w:tblCellSpacing", **{"w__w": 40, "w__type": "dxa"}))
        fixed_layout(tbl, widths)
        cell_margins(tbl, 40, 40, 60, 60)
        for i, (lab, colour) in enumerate(badges):
            cell = tbl.cell(0, i)
            shade(cell, colour)
            v_align(cell)
            p = cell.paragraphs[0]
            p.alignment = 1
            p.paragraph_format.space_after = self.Pt(0)
            self._run(p, lab, {"bold": True, "white": True, "size": 8.5})
        cell = tbl.cell(0, len(badges))
        v_align(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = self.Pt(0)
        for c in rest:
            if isinstance(c, str):
                t = norm(c.replace("\xa0", " "))
                if t:
                    self._run(p, (" " if p.text else "") + t, {})
            else:
                self._run(p, (" " if p.text and not p.text.endswith(" ") else ""), {})
                lab = norm(text_of(c))
                self._run(p, title_case(lab) if is_label(lab) else lab, {"bold": c.tag in ("b", "strong")})
        self.spacer(4)

    def tiles(self, n):
        kpis = [c for c in n.children if isinstance(c, Node)]
        if not kpis:
            return
        gap = 100
        w = (self.text_w - gap * (len(kpis) + 1)) // len(kpis)
        tbl = self.doc.add_table(rows=1, cols=len(kpis))
        no_borders(tbl)
        tbl._tbl.tblPr.append(_el("w:tblCellSpacing", **{"w__w": gap // 2, "w__type": "dxa"}))
        fixed_layout(tbl, [w] * len(kpis))
        tbl.alignment = 1
        cell_margins(tbl, 100, 100, 60, 60)
        for i, k in enumerate(kpis):
            cell = tbl.cell(0, i)
            shade(cell, "F8FAFC")
            tcPr = cell._tc.get_or_add_tcPr()
            bd = _el("w:tcBorders")
            badge = next((badge_of(c) for c in k.children if badge_of(c)), None)
            colour = badge[1] if badge else NAVY
            for s in ("left", "bottom", "right"):
                bd.append(_el(f"w:{s}", **{"w__val": "single", "w__sz": 4, "w__space": 0, "w__color": "D5DCE6"}))
            bd.append(_el("w:top", **{"w__val": "single", "w__sz": 24, "w__space": 0, "w__color": colour}))
            tcPr.append(bd)
            v_align(cell, "top")
            label = badge[0] if badge else ""
            number = next((norm(text_of(c)) for c in k.children if isinstance(c, Node) and c.tag == "b"), "")
            unit = norm("".join(c for c in k.children if isinstance(c, str)))
            p = cell.paragraphs[0]
            p.alignment = 1
            p.paragraph_format.space_after = self.Pt(2)
            self._run(p, label, {"bold": True, "colour": colour, "size": 9})
            p2 = cell.add_paragraph()
            p2.alignment = 1
            p2.paragraph_format.space_after = self.Pt(0)
            p2.paragraph_format.line_spacing = 1.0
            self._run(p2, number, {"bold": True, "size": 20, "colour": INK})
            p3 = cell.add_paragraph()
            p3.alignment = 1
            p3.paragraph_format.space_after = self.Pt(0)
            self._run(p3, title_case(unit) if unit else "", {"muted": True, "size": 8.5})
        self.spacer(8)

    def chart(self, n):
        bars = []
        rects = list(n.find_all("rect"))
        texts = [norm(text_of(t)) for t in n.find_all("text")]
        full = {"crit": "Critical", "high": "High", "medi": "Medium", "low": "Low", "info": "Info"}
        hmax = max([float(r.attrs.get("height") or 0) for r in rects] + [1])
        for i, r in enumerate(rects):
            lab = texts[i] if i < len(texts) else ""
            bars.append((full.get(lab, title_case(lab)), float(r.attrs.get("height") or 0) / hmax,
                         (r.attrs.get("fill") or "#475569").lstrip("#").upper()))
        if not bars:
            return
        cap = self.para("Meta")
        self._run(cap, title_case(n.attrs.get("aria-label") or "Chart"), {"bold": True, "muted": True})
        keep_next(cap)
        lw, bw = 1600, self.text_w - 1600
        tbl = self.doc.add_table(rows=len(bars), cols=2)
        no_borders(tbl)
        fixed_layout(tbl, [lw, bw])
        cell_margins(tbl, 30, 30, 60, 0)
        for i, (lab, frac, colour) in enumerate(bars):
            c0, c1 = tbl.cell(i, 0), tbl.cell(i, 1)
            p = c0.paragraphs[0]
            p.paragraph_format.space_after = self.Pt(0)
            self._run(p, lab, {"size": 9})
            bar = max(0, int(frac * (bw - 200) * 0.8))
            p = c1.paragraphs[0]
            p.paragraph_format.space_after = self.Pt(0)
            if bar < 60:
                self._run(p, "\u2014", {"muted": True, "size": 9})
                continue
            inner = c1.add_table(rows=1, cols=1)
            no_borders(inner)
            fixed_layout(inner, [bar])
            shade(inner.cell(0, 0), colour)
            ip = inner.cell(0, 0).paragraphs[0]
            ip.paragraph_format.space_after = self.Pt(0)
            self._run(ip, "\u00a0", {"size": 7})
            p._p.getparent().remove(p._p)
            tail = c1.add_paragraph()
            tail.paragraph_format.space_after = self.Pt(0)
            tail.paragraph_format.line_spacing = self.Pt(1)
        self.spacer(8)

    def table(self, n):
        rows = []
        for tr in n.find_all("tr"):
            cells = [c for c in tr.children if isinstance(c, Node) and c.tag in ("td", "th")]
            if cells:
                rows.append(cells)
        if not rows:
            return
        ncols = max(sum(int(c.attrs.get("colspan") or 1) for c in r) for r in rows)
        header = all(c.tag == "th" for c in rows[0]) and len(rows[0]) == ncols and ncols > 1
        widths = self._widths(rows, ncols)
        tbl = self.doc.add_table(rows=len(rows), cols=ncols)
        table_borders(tbl)
        fixed_layout(tbl, widths)
        cell_margins(tbl)
        tbl.alignment = 0
        for ri, r in enumerate(rows):
            row = tbl.rows[ri]
            cant_split(row)
            ci = 0
            for c in r:
                span = int(c.attrs.get("colspan") or 1)
                cell = tbl.cell(ri, ci)
                if span > 1 and ci + span <= ncols:
                    cell = cell.merge(tbl.cell(ri, ci + span - 1))
                is_head = ri == 0 and header
                self.cell(cell, c, head=is_head, key=(c.tag == "th" and not is_head),
                          zebra=(not is_head and ri % 2 == 0 and header))
                ci += span
            if ri == 0 and header:
                repeat_header(row)
        self.spacer(8)

    def _widths(self, rows, ncols):
        """Column widths in twips: content-driven, HTML width hints act as caps, badges get room."""
        hint = [0] * ncols
        avg = [[] for _ in range(ncols)]
        word = [3] * ncols
        badge = [False] * ncols
        for r in rows:
            ci = 0
            for c in r:
                if ci >= ncols:
                    break
                span = int(c.attrs.get("colspan") or 1)
                m = re.search(r"width:\s*(\d+)px", c.attrs.get("style", ""))
                if m:
                    hint[ci] = int(m.group(1)) * 15
                if span == 1:
                    t = norm(text_of(c))
                    avg[ci].append(len(t))
                    word[ci] = max(word[ci], min(max((len(w) for w in t.split()), default=0), 28))
                    if any(badge_of(x) for x in c.children):
                        badge[ci] = True
                ci += span
        floor, want = [], []
        for i in range(ncols):
            a = sum(avg[i]) / max(1, len(avg[i]))
            f = max(min(word[i], 12) * 95 + 220, 1650 if badge[i] else 700)
            nat = max(f, min(a, 140) * 105 + 260)
            if hint[i]:
                nat = max(f, min(nat, hint[i]))
            floor.append(f)
            want.append(nat)
        fixed = [False] * ncols
        for _ in range(ncols):  # proportional fit that never squeezes a column below its longest word
            free = self.text_w - sum(f for f, x in zip(floor, fixed) if x)
            flex = sum(w for w, x in zip(want, fixed) if not x) or 1
            changed = False
            for i in range(ncols):
                if not fixed[i] and want[i] * free / flex < floor[i]:
                    fixed[i], want[i], changed = True, floor[i], True
            if not changed:
                want = [w if x else w * free / flex for w, x in zip(want, fixed)]
                break
        total = sum(want)
        w = [int(x * self.text_w / total) for x in want]
        w[-1] += self.text_w - sum(w)
        return w

    def cell(self, cell, node, *, head=False, key=False, zebra=False):
        v_align(cell, "center")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = self.Pt(0)
        p.paragraph_format.line_spacing = 1.05
        kids = [c for c in node.children if not (isinstance(c, str) and not c.strip())]
        first_badge = badge_of(kids[0]) if kids else None
        if head:
            shade(cell, HEAD_FILL)
            self._run(p, title_case(norm(text_of(node))), {"bold": True, "white": True, "size": 9})
            return
        if key:
            shade(cell, KEY_FILL)
            self._run(p, title_case(norm(text_of(node))), {"bold": True, "size": 9.5})
            return
        if first_badge:
            lab, colour = first_badge
            shade(cell, colour)
            p.alignment = 1
            self._run(p, lab, {"bold": True, "white": True, "size": 9})
            rest = kids[1:]
            txt = norm("".join(text_of(c) for c in rest if not (isinstance(c, Node) and c.tag == "br")))
            if txt:
                p2 = cell.add_paragraph()
                p2.alignment = 1
                p2.paragraph_format.space_after = self.Pt(0)
                self._run(p2, title_case(txt) if is_label(txt) else txt, {"white": True, "size": 8, "italic": True})
            return
        if zebra:
            shade(cell, ZEBRA)
        whole = norm(text_of(node))
        if whole in ("-", "\u2014", "") or re.fullmatch(r"[\d.,/%]+", whole or "x"):
            p.alignment = 1
        has_block = any(isinstance(c, Node) and (c.tag in ("table", "pre", "ul", "ol", "p", "div")) for c in kids)
        if not has_block:
            for c in node.children:
                self.runs(p, c, {"size": 9.5}, label=True)
            return
        # complex cell: render nested blocks inside the cell
        for c in node.children:
            if isinstance(c, Node) and c.tag in ("p", "div", "pre", "ul", "ol"):
                p = cell.add_paragraph()
                p.paragraph_format.space_after = self.Pt(2)
                self.runs(p, c, {"size": 9.5, "mono": c.tag == "pre"})
            elif isinstance(c, Node) and c.tag == "table":
                p = cell.add_paragraph()
                self._run(p, norm(text_of(c)), {"size": 9})
            else:
                self.runs(p, c, {"size": 9.5}, label=True)
        if not cell.paragraphs[0].text and len(cell.paragraphs) > 1:
            first = cell.paragraphs[0]._p
            first.getparent().remove(first)

    # ---- page furniture
    def cover(self, cover: Node | None, title: str, tenant: str, letter: bool):
        if letter or cover is None:
            return
        for _ in range(5):
            self.spacer(18)
        metas, subtitle = [], ""
        brand_line = ""
        for c in cover.children:
            if not isinstance(c, Node):
                continue
            t = norm(text_of(c))
            if "meta" in c.classes and not brand_line and c.tag == "div":
                brand_line = t
            elif c.tag == "h1":
                pass
            elif "meta" in c.classes:
                metas.append(t)
            elif c.tag == "p" and not subtitle:
                subtitle = t
        p = self.para()
        self._run(p, "ATTESTRA", {"bold": True, "colour": NAVY, "size": 12})
        rest = brand_line.replace("ATTESTRA", "").strip(" ·")
        if rest:
            self._run(p, f"   \u00b7   {rest}", {"muted": True, "size": 11})
        para_border_bottom(p, NAVY, 12)
        p.paragraph_format.space_after = self.Pt(24)
        t = self.doc.add_paragraph(title_case(title), style="Title")
        t.paragraph_format.space_after = self.Pt(10)
        if subtitle:
            sp = self.para()
            self._run(sp, subtitle, {"size": 16, "colour": INK})
            sp.paragraph_format.space_after = self.Pt(30)
        tbl = self.doc.add_table(rows=0, cols=2)
        no_borders(tbl)
        rows = []
        for m in metas:
            for part in re.split(r"\s+\u00b7\s+", m):
                k, _, v = part.partition(":")
                if v:
                    rows.append((k.strip(), v.strip()))
                else:
                    m2 = re.match(r"^(.*?\D)\s+(#?\d[\w.-]*)$", part.strip())
                    rows.append((m2.group(1), m2.group(2)) if m2 else ("", part.strip()))
        for k, v in rows:
            r = tbl.add_row()
            self._run(r.cells[0].paragraphs[0], title_case(k), {"bold": True, "muted": True, "size": 10})
            self._run(r.cells[1].paragraphs[0], v, {"size": 10})
        if rows:
            fixed_layout(tbl, [2600, self.text_w - 2600])
            cell_margins(tbl, 50, 50, 0, 80)
        for _ in range(4):
            self.spacer(18)
        status = self.para()
        if self.draft:
            self._run(status, "\u00a0DRAFT \u2014 NOT SIGNED\u00a0", {"bold": True, "white": True, "size": 10})
            run_shade(status.runs[-1], "A1161B")
        else:
            self._run(status, "\u00a0SIGNED\u00a0", {"bold": True, "white": True, "size": 10})
            run_shade(status.runs[-1], "15803D")
        conf = self.para("Meta")
        self._run(conf, "Confidential. Prepared for the named client only.", {"muted": True})
        self.doc.add_page_break()

    def toc(self, entries: list, depth: int = 2):
        h = self.para()
        self._run(h, "Contents", {"bold": True, "size": 16, "colour": NAVY})
        para_border_bottom(h, NAVY, 8)
        h.paragraph_format.space_after = self.Pt(10)
        p = self.para()
        r = p.add_run()
        r._r.append(_el("w:fldChar", **{"w__fldCharType": "begin", "w__dirty": "true"}))
        r2 = p.add_run()
        it = _el("w:instrText")
        it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        it.text = f' TOC \\o "1-{depth}" \\h \\z \\u '
        r2._r.append(it)
        r3 = p.add_run()
        r3._r.append(_el("w:fldChar", **{"w__fldCharType": "separate"}))
        # cached result: the headings, so the page reads correctly before Word refreshes the field
        first = True
        for lvl, txt in entries:
            q = p if first else self.para(f"TOC {lvl}")
            if first:
                p.style = self.doc.styles[f"TOC {lvl}"]
            self._run(q, txt, {})
            first = False
        end_p = p if first else self.doc.paragraphs[-1]
        r5 = end_p.add_run()
        r5._r.append(_el("w:fldChar", **{"w__fldCharType": "end"}))
        note = self.para("Small")
        note.paragraph_format.space_before = self.Pt(10)
        self._run(note, "Page numbers in this contents list are filled in by Word (References \u203a Update Table).", {})
        self.doc.add_page_break()
        # ask Word to refresh fields (TOC page numbers) on open
        settings = self.doc.settings.element
        settings.append(_el("w:updateFields", **{"w__val": "true"}))

    def furniture(self, title: str, tenant: str, letter: bool):
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = not letter
        hdr = sec.header
        p = hdr.paragraphs[0]
        p.alignment = 2
        self._run(p, tenant + "  \u00b7  " if tenant else "", {"muted": True, "size": 8.5})
        self._run(p, title_case(title), {"bold": True, "colour": NAVY, "size": 8.5})
        para_border_bottom(p, "C9D1DC", 4)
        if self.draft:
            from lxml import etree
            band = hdr.add_paragraph()
            band.alignment = 1
            para_shade(band, "FDECEC")
            band.paragraph_format.space_before = self.Pt(3)
            self._run(band, "DRAFT \u2014 not signed. Not for distribution until signed in Attestra.",
                      {"bold": True, "colour": "A1161B", "size": 8})
            for hp in ([p] + ([sec.first_page_header.paragraphs[0]] if not letter else [])):
                hp._p.append(etree.fromstring(WATERMARK_XML))
        ftr = sec.footer.paragraphs[0]
        ftr.alignment = 1
        para_top = ftr._p.get_or_add_pPr()
        bdr = _el("w:pBdr")
        bdr.append(_el("w:top", **{"w__val": "single", "w__sz": 4, "w__space": 4, "w__color": "C9D1DC"}))
        para_top.append(bdr)
        self._run(ftr, "Confidential  \u00b7  ", {"muted": True, "size": 8.5})
        self._run(ftr, "Page ", {"muted": True, "size": 8.5})
        for r in add_field(ftr, "PAGE"):
            r.font.size, r.font.color.rgb = self.Pt(8.5), self.RGB.from_string(MUTED)
        self._run(ftr, " of ", {"muted": True, "size": 8.5})
        for r in add_field(ftr, "NUMPAGES"):
            r.font.size, r.font.color.rgb = self.RGB and self.Pt(8.5), self.RGB.from_string(MUTED)


def _find(node, pred):
    for c in node.children:
        if isinstance(c, Node):
            if pred(c):
                return c
            r = _find(c, pred)
            if r is not None:
                return r
    return None


def render_docx(html: str, *, kind: str = "", draft: bool = True, base_url: str | None = None) -> bytes:
    """Rebuild the frozen HTML report as a styled Word document and return its bytes."""
    try:  # same display rules as the PDF: framework names, IST times (display only)
        from attestra.reports.pdf import _framework_names, _human_times

        html = _framework_names(_human_times(html))
    except Exception:
        pass
    root = parse(html)
    body = _find(root, lambda n: n.tag == "body") or root
    cover = _find(body, lambda n: n.tag == "section" and "cover" in n.classes)
    letter = kind in LETTER_KINDS or cover is None
    title_node = (cover and _find(cover, lambda n: n.tag == "h1")) or _find(body, lambda n: n.tag == "h1")
    title = norm(text_of(title_node)) if title_node else (kind.replace("_", " ").title() or "Report")
    brand = _find(body, lambda n: "meta" in n.classes and _find(n, lambda m: "brand" in m.classes) is not None)
    tenant = ""
    if brand is not None:
        parts = [x.strip() for x in norm(text_of(brand)).split("\u00b7")]
        tenant = parts[1] if len(parts) > 1 else ""

    b = Builder(kind=kind, draft=draft, base_url=base_url)
    body_h1 = [h for h in body.find_all("h1") if not (cover is not None and h in list(cover.find_all("h1")))]
    b.h_offset = 0 if body_h1 else 1

    # pre-compute headings for the cached TOC
    entries = []
    for h in _iter_headings(body, cover):
        lvl = int(h.tag[1]) - b.h_offset
        if 1 <= lvl <= 2:
            t = norm(text_of(h))
            entries.append((lvl, (display(t) if re.match(r"^F-\d+", t) else title_case(display(t)))))

    b.furniture(title, tenant, letter)
    b.cover(cover, title, tenant, letter)
    depth = 2 if len(entries) <= 25 else 1
    entries = [e for e in entries if e[0] <= depth]
    if not letter and len(entries) > 2:
        b.toc(entries, depth)
    b.blocks(body)
    core = b.doc.core_properties
    core.title, core.subject, core.author = title_case(title), kind or "report", "Attestra"
    core.keywords = "DRAFT" if draft else "SIGNED"
    out = BytesIO()
    b.doc.save(out)
    return out.getvalue()


def _iter_headings(node, cover):
    for c in node.children:
        if not isinstance(c, Node) or c is cover or c.tag in SKIP or (c.tag == "nav"):
            continue
        if c.tag in ("h1", "h2", "h3"):
            yield c
        else:
            yield from _iter_headings(c, cover)


def docx_text(data: bytes) -> str:
    """All visible text of a .docx (body, tables, headers, footers) for coverage checks."""
    import zipfile
    with zipfile.ZipFile(BytesIO(data)) as z:
        xml = " ".join(z.read(n).decode("utf-8", "replace") for n in z.namelist()
                       if n.startswith("word/") and n.endswith(".xml"))
    xml = re.sub(r"</w:p>|<w:br/>|<w:tab/>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml)
