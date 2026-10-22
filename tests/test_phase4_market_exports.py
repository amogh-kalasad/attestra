"""Phase 4 market-facing export checks: PDF engine, OSCAL, Word, JSON, XLSX."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from attestra.models import AuditEvent
from attestra.reports.render import KINDS

FRAMEWORK_FOR = {"compliance": "pci_dss_4", "iso_soa": "iso_27001_2022"}


# @lat: [[tests#Phase 4 market exports#PDF engine status probe]]
def test_pdf_engine_status_shape():
    from attestra.reports.pdf import pdf_engine_status

    status = pdf_engine_status()
    assert set(status) >= {"available", "engine", "message"}
    assert isinstance(status["available"], bool)
    assert status["engine"] in ("weasyprint", "xhtml2pdf", "none")
    assert isinstance(status["message"], str) and status["message"]
    if status["available"]:
        assert status["engine"] != "none"
        assert "available" in status["message"].lower()
    else:
        assert status["engine"] == "none"
        assert "missing" in status["message"].lower()
        assert "xhtml2pdf" in status["message"] and "weasyprint" in status["message"]


# @lat: [[tests#Phase 4 market exports#OSCAL AR and POA&M schema validation]]
def test_oscal_ar_and_poam_schema_validation(loaded):
    from attestra import db
    from attestra.reports.oscal import (
        assessment_results,
        load_oscal_schema,
        poam,
        schema_for_oscal_kind,
        validate_oscal_structure,
    )

    schema = load_oscal_schema()
    assert "required_top_level_poam" in schema
    ar_schema = schema_for_oscal_kind(schema, "assessment-results")
    poam_schema = schema_for_oscal_kind(schema, "poam")

    with db.session() as s:
        ar = assessment_results(s, loaded, "pci_dss_4")
        p = poam(s, loaded)

    assert validate_oscal_structure(ar, ar_schema) == []
    assert validate_oscal_structure(p, poam_schema) == []

    broken_ar = {"assessment-results": {"uuid": "broken"}}
    ar_errors = validate_oscal_structure(broken_ar, ar_schema)
    assert ar_errors
    assert any("metadata" in e or "results" in e for e in ar_errors)

    broken_poam = {"plan-of-action-and-milestones": {"uuid": "broken"}}
    poam_errors = validate_oscal_structure(broken_poam, poam_schema)
    assert poam_errors
    assert any("metadata" in e or "poam-items" in e for e in poam_errors)

    assert validate_oscal_structure({}, ar_schema)
    assert validate_oscal_structure({"assessment-results": "not-an-object"}, ar_schema)


# @lat: [[tests#Phase 4 market exports#Word re-export twice every kind]]
@pytest.mark.parametrize("kind", KINDS)
def test_word_reexport_twice_every_kind(loaded, kind):
    from attestra import db
    from attestra.reports import generate_report
    from attestra.reports.export import export_report_formats

    with db.session() as s:
        r = generate_report(s, loaded, kind=kind, framework=FRAMEWORK_FOR.get(kind, ""), actor="tester")
        first = export_report_formats(r)
        second = export_report_formats(r)
        docx1 = Path(first["docx_path"]).read_bytes()
        docx2 = Path(second["docx_path"]).read_bytes()
    assert docx1[:2] == b"PK"
    assert docx2[:2] == b"PK"
    assert first["docx_path"] == second["docx_path"]


# @lat: [[tests#Phase 4 market exports#JSON engagement export shape]]
def test_engagement_json_export_shape_and_audit(loaded):
    """Version fields, top-level keys, controls/reports present; audited as engagement.exported."""
    import json

    from sqlalchemy import select

    from attestra import __version__, db, services
    from attestra.models import AuditEvent
    from attestra.reports import generate_report
    from attestra.reports.engagement_export import EXPORT_SCHEMA_VERSION, export_engagement_json

    with db.session() as s:
        findings = services.list_findings(s, loaded)
        assert findings
        services.add_evidence(
            s,
            findings[0].id,
            kind="note",
            filename="note.txt",
            data=b"validation note for export",
            note="meta note",
            actor="tester",
        )
        generate_report(s, loaded, kind="technical", actor="tester")
        payload = export_engagement_json(s, loaded, actor="exporter@test")

        assert payload["version"] == __version__
        assert payload["export_schema_version"] == EXPORT_SCHEMA_VERSION == 1
        assert set(payload) >= {"version", "export_schema_version", "engagement", "findings", "controls", "evidence", "reports"}

        eng = payload["engagement"]
        assert eng["id"] == loaded
        assert eng["name"]
        assert eng["type"] == "vapt"
        assert eng["type_label"] == "VA + PT"
        assert "pci_dss_4" in eng["frameworks"]
        assert eng["tenant"]
        assert eng["window_start"] and eng["window_end"]

        assert isinstance(payload["findings"], list) and payload["findings"]
        sample = payload["findings"][0]
        for key in ("id", "title", "severity", "status", "impact", "affected_urls", "references", "cvss_vector", "location"):
            assert key in sample

        assert isinstance(payload["controls"], dict)
        assert "pci_dss_4" in payload["controls"]
        assert isinstance(payload["controls"]["pci_dss_4"], list)
        assert payload["controls"]["pci_dss_4"]
        assert "ref" in payload["controls"]["pci_dss_4"][0]
        assert "status" in payload["controls"]["pci_dss_4"][0]

        assert isinstance(payload["evidence"], list) and payload["evidence"]
        assert isinstance(payload["reports"], list) and payload["reports"]
        report_row = payload["reports"][0]
        for key in ("id", "kind", "framework", "version", "status", "sha256", "generated_at", "generated_by"):
            assert key in report_row
        assert "path" not in report_row

        # Must be JSON-serializable without a custom default.
        json.dumps(payload)

        events = list(
            s.scalars(
                select(AuditEvent).where(
                    AuditEvent.action == "engagement.exported",
                    AuditEvent.entity == "engagement",
                    AuditEvent.entity_id == loaded,
                )
            )
        )
        assert events
        assert events[-1].details.get("format") == "json"


# @lat: [[tests#Phase 4 market exports#JSON engagement export evidence metadata only]]
def test_engagement_json_export_evidence_no_body_keys(loaded):
    """Evidence rows expose only metadata keys — never body/path/bytes/content."""
    from attestra import db, services
    from attestra.reports.engagement_export import export_engagement_json

    allowed = {"kind", "filename", "sha256", "note", "finding_id"}
    forbidden = {"body", "bytes", "content", "data", "path", "password", "secret", "api_key", "token"}

    with db.session() as s:
        findings = services.list_findings(s, loaded)
        services.add_evidence(
            s,
            findings[0].id,
            kind="note",
            filename="cap.txt",
            data=b"raw capture body must not appear",
            note="ok note",
            actor="tester",
        )
        payload = export_engagement_json(s, loaded, actor="tester")

    assert payload["evidence"]
    for row in payload["evidence"]:
        assert set(row.keys()) <= allowed
        assert not (forbidden & set(row.keys()))
        assert "raw capture body" not in str(row)


# @lat: [[tests#Phase 4 market exports#JSON engagement export formula and secret absence]]
def test_engagement_json_export_formula_and_secret_absence(loaded):
    """Evidence bodies with secrets or formula-like text must not leak into the JSON export."""
    import json

    from attestra import db, services
    from attestra.reports.engagement_export import export_engagement_json

    secret = "s3cret-export-token-9f3a"
    formula = "=cmd|'/c calc'!A0"
    body = f"password={secret}\napi_key={secret}\n{formula}\n"

    with db.session() as s:
        findings = services.list_findings(s, loaded)
        f = findings[0]
        f.title = formula  # formula-like title stays a plain string in JSON (not an Excel cell)
        services.add_evidence(
            s,
            f.id,
            kind="log",
            filename="leak.txt",
            data=body.encode(),
            note="masked capture",
            actor="tester",
        )
        payload = export_engagement_json(s, loaded, actor="tester")
        dumped = json.dumps(payload)

    # Secret from the evidence *body* must not appear (bodies are excluded; masking may also redact).
    assert secret not in dumped
    assert "password=" not in dumped
    # Evidence metadata only — no body field to carry the formula blob.
    for row in payload["evidence"]:
        assert "body" not in row
        assert formula not in (row.get("note") or "")
    # Finding title may contain formula-like text as data; export must remain plain JSON (no CSV cell).
    titles = [f["title"] for f in payload["findings"]]
    assert formula in titles
    assert "'" + formula not in dumped  # CSV-style quote-prefix must not be applied here


# @lat: [[tests#Phase 4 market exports#XLSX three sheets]]
def test_xlsx_three_sheets(loaded):
    """Workbook has Findings, Controls, and Evidence sheets with headers and freeze/filter."""
    from attestra import db
    from attestra.reports.findings_csv import CSV_COLUMNS
    from attestra.reports.xlsx_export import (
        CONTROLS_COLUMNS,
        EVIDENCE_COLUMNS,
        export_engagement_xlsx,
    )

    with db.session() as s:
        data = export_engagement_xlsx(s, loaded, actor="tester")

    assert data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(data), read_only=False)
    assert wb.sheetnames == ["Findings", "Controls", "Evidence"]

    for name, expected_headers in (
        ("Findings", list(CSV_COLUMNS)),
        ("Controls", list(CONTROLS_COLUMNS)),
        ("Evidence", list(EVIDENCE_COLUMNS)),
    ):
        ws = wb[name]
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert headers == expected_headers
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref.startswith("A1:")

    assert wb["Findings"].max_row >= 2
    assert wb["Controls"].max_row >= 2


# @lat: [[tests#Phase 4 market exports#XLSX formula guard]]
def test_xlsx_formula_guard(loaded):
    """A finding title of =cmd() is prefixed with a single quote in the Findings sheet."""
    from attestra import db, services
    from attestra.reports.xlsx_export import export_engagement_xlsx

    with db.session() as s:
        findings = services.list_findings(s, loaded)
        assert findings
        findings[0].title = "=cmd()"
        data = export_engagement_xlsx(s, loaded, actor="tester")

    wb = load_workbook(io.BytesIO(data))
    ws = wb["Findings"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    title_idx = headers.index("title")
    titles = [row[title_idx].value for row in ws.iter_rows(min_row=2) if row[title_idx].value]
    assert any(t == "'=cmd()" for t in titles)
    assert not any(isinstance(t, str) and t.startswith("=") for t in titles)


# @lat: [[tests#Phase 4 market exports#XLSX evidence metadata only]]
def test_xlsx_evidence_metadata_only(loaded):
    """Evidence sheet lists metadata only; body bytes and secrets never appear in the workbook."""
    from attestra import db, services
    from attestra.reports.xlsx_export import export_engagement_xlsx

    secret = "SUPER_SECRET_EVIDENCE_BODY_9f3a"
    with db.session() as s:
        findings = services.list_findings(s, loaded)
        assert findings
        ev = services.add_evidence(
            s,
            findings[0].id,
            kind="note",
            filename="cap.txt",
            data=secret.encode(),
            note="capture note",
            actor="tester",
        )
        ev_id = ev.id
        finding_id = findings[0].id
        sha = ev.sha256
        data = export_engagement_xlsx(s, loaded, actor="tester")

    assert secret.encode() not in data
    assert secret not in data.decode("latin-1", errors="ignore")

    wb = load_workbook(io.BytesIO(data))
    ws = wb["Evidence"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "path" not in headers
    assert headers == [
        "finding_id",
        "evidence id",
        "kind",
        "filename",
        "sha256",
        "note",
    ]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    match = [r for r in rows if str(r[1]) == str(ev_id)]
    assert match
    row = match[0]
    assert str(row[0]) == str(finding_id)
    assert row[2] == "note"
    assert row[3] == "cap.txt"
    assert row[4] == sha
    assert row[5] == "capture note"
    for cell in row:
        assert secret not in str(cell)


# @lat: [[tests#Phase 4 market exports#XLSX export audit]]
def test_xlsx_export_audits_engagement_exported(loaded):
    """export_engagement_xlsx records engagement.exported with format=xlsx."""
    from attestra import db
    from attestra.reports.xlsx_export import export_engagement_xlsx

    with db.session() as s:
        export_engagement_xlsx(s, loaded, actor="tester")
        events = list(
            s.scalars(
                select(AuditEvent).where(
                    AuditEvent.action == "engagement.exported",
                    AuditEvent.entity == "engagement",
                    AuditEvent.entity_id == loaded,
                )
            )
        )
    assert events
    assert events[-1].details.get("format") == "xlsx"
    assert events[-1].actor == "tester"


# @lat: [[tests#Phase 4 market exports#Health PDF engine fields]]
def test_health_includes_pdf_engine_fields(client):
    """GET /api/health exposes pdf probe fields with a clear message when engines are missing."""
    r = client.get("/api/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "ok"
    assert "version" in body
    assert body["pdf_engine"] in ("weasyprint", "xhtml2pdf", "none")
    assert body["pdf_message"]
    assert set(body["pdf"]) >= {"available", "engine", "message"}
    if not body["pdf"]["available"]:
        assert "missing" in body["pdf_message"].lower()
        assert "xhtml2pdf" in body["pdf_message"]


# @lat: [[tests#Phase 4 market exports#API and UI Excel JSON routes]]
def test_api_and_ui_xlsx_json_routes(client, loaded, auth_headers):
    """API and engagement page routes return Excel and JSON engagement exports."""
    api_x = client.get(f"/api/engagements/{loaded}/export/xlsx", headers=auth_headers)
    assert api_x.status_code == 200, api_x.text
    assert api_x.content[:2] == b"PK"
    assert "spreadsheetml" in api_x.headers.get("content-type", "")

    api_j = client.get(f"/api/engagements/{loaded}/export/json", headers=auth_headers)
    assert api_j.status_code == 200, api_j.text
    payload = api_j.json()
    assert payload["export_schema_version"] == 1
    assert "findings" in payload and "controls" in payload

    client.headers.pop("Authorization", None)
    assert client.post(
        "/login", data={"username": "lead", "password": "test-pass-lead"}, follow_redirects=False
    ).status_code == 303
    page = client.get(f"/engagements/{loaded}")
    assert f"/engagements/{loaded}/export/xlsx" in page.text
    assert f"/engagements/{loaded}/export/json" in page.text

    ui_x = client.get(f"/engagements/{loaded}/export/xlsx")
    assert ui_x.status_code == 200 and ui_x.content[:2] == b"PK"
    ui_j = client.get(f"/engagements/{loaded}/export/json")
    assert ui_j.status_code == 200
    assert ui_j.json()["version"]
